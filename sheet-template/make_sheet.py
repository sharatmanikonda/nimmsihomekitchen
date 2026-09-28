"""Builds nimmi-menu.xlsx, the starting point for Nimmi's menu Google Sheet.

Upload the file to Google Drive and open it with Google Sheets (or File -> Import), share it as
"Anyone with the link -> Viewer", and put its id in menuSheetId in src/data/kitchen.js.

    pip install openpyxl
    python sheet-template/make_sheet.py
"""
import datetime
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
src = open(os.path.join(ROOT, 'src', 'data', 'kitchen.js'), encoding='utf-8').read()
catalogue = re.findall(r"\{ id: '[^']+', cat: '([^']+)', name: '([^']+)'.*?price: (\d+), unit: '([^']+)'(, nonVeg: true)?", src)
cat_label = dict(re.findall(r"\{ id: '([^']+)', label: '([^']+)'", src))

# Everyday dishes first (from Nimmi's WhatsApp posts), then the catalogue.
DAILY = 'Lunch / dinner'
dishes = [
    ('Sambar', '500g', '85', 'Veg', DAILY),
    ('Brinjal fry', '150g', '75', 'Veg', DAILY),
] + [(name, unit, price, 'Non-veg' if nv else 'Veg', cat_label[cat]) for cat, name, price, unit, nv in catalogue]

HEAD_FONT = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2F4A2A')


def header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = 'A2'


def dropdown(ws, options, cells):
    dv = DataValidation(type='list', formula1='"' + ','.join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)


wb = Workbook()
menu = wb.active
menu.title = 'Menu'
# Portion, Price and Limit are stored as text so "85/-" or "₹85" typed later stay in the same
# column type; the Sheets query API drops cells whose type differs from the rest of the column.
menu.append(['Dish', 'Portion', 'Price', 'Today', 'Limit', 'Type', 'Category'])
for name, unit, price, kind, cat in dishes:
    menu.append([name, unit, price, None, None, kind, cat])
last = 60  # room for new dishes
header(menu, [38, 18, 10, 10, 10, 11, 20])
for col in 'BCE':
    for r in range(2, last + 1):
        menu[f'{col}{r}'].number_format = '@'
dropdown(menu, ['Yes', 'No'], f'D2:D{last}')
dropdown(menu, ['Veg', 'Non-veg'], f'F2:F{last}')
dropdown(menu, [DAILY] + list(cat_label.values()), f'G2:G{last}')

st = wb.create_sheet('Settings')
st.append(['Delivery date', 'Cutoff time', 'Cutoff date', 'Delivery slots', 'Note'])
st.append([None, datetime.time(10, 30), None, 'Lunch', None])
header(st, [16, 14, 16, 40, 50])
st['A2'].number_format = st['C2'].number_format = 'dd/mm/yyyy'
st['B2'].number_format = 'h:mm AM/PM'
dates = DataValidation(type='date', operator='greaterThan', formula1='1', allow_blank=True)
st.add_data_validation(dates)
dates.add('A2')
dates.add('C2')
for c in 'ABCDE':
    st[f'{c}2'].alignment = Alignment(wrap_text=True, vertical='top')

how = wb.create_sheet('How to post')
for line in [
    "How to post today's menu",
    '',
    '1. Settings tab: pick the Delivery date (A2). Change the Cutoff time if needed (orders close at that time on the delivery day).',
    '2. Menu tab: set Today to "Yes" for each dish you are making, and "No" for the rest.',
    '   Check the Portion and Price. For a new dish, add a row at the bottom (Dish, Portion, Price, Today = Yes).',
    '3. Optional: put a number in Limit to cap how many can be ordered. Add a Note in Settings for customers.',
    '4. Open the website\'s "Post today\'s menu" page, check the menu, and tap "Share to WhatsApp group".',
    '',
    'The website shows the new menu within a minute or two.',
    'After the cutoff, the website shows "menu coming soon" until you set a new Delivery date.',
    'Cutoff date: leave empty for the same day as delivery; fill it in only if orders close on an earlier day.',
    '',
    'Please do not rename the tabs or change the row 1 headings.',
]:
    how.append([line])
how['A1'].font = Font(bold=True, size=14)
how.column_dimensions['A'].width = 120

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'nimmi-menu.xlsx')
wb.save(out)
print(f'{len(dishes)} dishes -> {out}')
